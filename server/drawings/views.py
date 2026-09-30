import io

from django.core.paginator import Paginator
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, render
from django.urls import reverse
from django.conf import settings
from pathlib import Path

from .models import AuditLog, BomItem, Drawing
from . import access, browse, relations
from .sources import open_source


PER_PAGE = 50


def _int(v, default=None):
    try:
        return max(0, int(v))
    except (TypeError, ValueError):
        return default


def _page_window(page, around=2):
    """ページ番号の並び（間は None＝「…」）。"""
    n, cur = page.paginator.num_pages, page.number
    keep = {1, n, *range(cur - around, cur + around + 1)}
    out, prev = [], 0
    for i in sorted(x for x in keep if 1 <= x <= n):
        if i - prev > 1:
            out.append(None)
        out.append(i)
        prev = i
    return out


def _state(request):
    """見せ方は URL で指定がなければ前回使ったもの（セッションに覚える）。"""
    st = browse.state_from(request.GET, request.session.get("view", "list"), access.can_see_all(request.user))
    if request.GET.get("v") in browse.VIEWS:
        request.session["view"] = st["v"]
    return st


def index(request):
    st = _state(request)
    terms, qs = browse.filtered(st, request.user)
    page = Paginator(browse.groups(qs, st["v"]), PER_PAGE).get_page(request.GET.get("page"))
    base = page.start_index() - 1 if page.paginator.count else 0
    items = [{"i": base + n, "row": r, "d": d, "size": size, "link": browse.qs_string(st, i=base + n)}
             for n, (r, d, size) in enumerate(browse.drawings_for(page.object_list))]
    tabs = [(k, label, browse.qs_string({"q": st["q"], "kind": st["kind"], "v": k}) or "?v=list")
            for k, label in browse.VIEWS.items()]
    # 階層は、検索語・種類だけで絞った全体から作る（選んだ枝の兄弟も見えるように）
    nodes = browse.tree(st, browse.filtered({**st, "f": "", "m": "", "s": ""}, request.user)[1]) if st["v"] != "list" else []
    for nd in nodes:
        key = {"folder": "f", "model": "m", "series": "s"}[st["v"]]
        nd["link"] = browse.qs_string(st, **{key: nd["value"]}) if nd["value"] is not None else ""
    sel_key = {"folder": "f", "model": "m", "series": "s"}.get(st["v"])
    crumbs = [(label, browse.qs_string(st, **{sel_key: val})) for label, val in browse.crumbs(st)] if sel_key else []
    return render(request, "drawings/search.html", {
        "st": st, "q": st["q"], "kind": st["kind"], "terms": terms, "page": page, "total": page.paginator.count,
        "drawing_total": qs.count() if st["v"] != "list" or st["q"] else None,
        "items": items, "window": _page_window(page), "page_qs": browse.qs_string(st),
        "tabs": tabs, "nodes": nodes, "crumbs": crumbs,
        "top_link": browse.qs_string(st, **{sel_key: ""}) if sel_key else "", "view_label": browse.VIEWS[st["v"]],
        "see_all": access.can_see_all(request.user), "doc_filters": browse.DOC_FILTERS,
    })


def detail(request, pk):
    """一覧から開いたときは一覧の条件と i（一覧での位置）を受け取り、
    「一覧に戻る」で同じページ・同じ位置へ戻れるようにし、前後の図面（ファイル）へ移れるようにする。"""
    d = get_object_or_404(access.visible(request.user, Drawing.objects.select_related("file")), pk=pk)
    AuditLog.objects.create(user=request.user, action=AuditLog.Action.VIEW, target=d.file.path)
    st = browse.state_from(request.GET, request.session.get("view", "list"), access.can_see_all(request.user))
    i = _int(request.GET.get("i"))
    back_page = _int(request.GET.get("p"), 1) or 1
    nav = {}
    if i is not None:
        back_page = i // PER_PAGE + 1
        _, qs = browse.filtered(st, request.user)
        ordered = browse.groups(qs, st["v"])
        lo = max(i - 1, 0)
        rows = browse.drawings_for(ordered[lo:i + 2])
        cur = i - lo
        if cur < len(rows) and rows[cur][1].file_id == d.file_id:  # 一覧が変わっていなければ前後を出す
            nav["total"] = ordered.count()
            nav["pos"] = i + 1
            if cur > 0:
                nav["prev"] = (rows[cur - 1][1], browse.qs_string(st, i=i - 1))
            if cur + 1 < len(rows):
                nav["next"] = (rows[cur + 1][1], browse.qs_string(st, i=i + 1))
    back = reverse("index") + (browse.qs_string(st, page=back_page if back_page > 1 else "") or "?v=" + st["v"])
    if i is not None:
        back += f"#f{d.file_id}"
    keep = browse.qs_string(st, p=back_page if back_page > 1 else "")  # 同じファイルの図面などへ移っても戻り先を保つ
    revisions = (access.visible(request.user).filter(drawing_no=d.drawing_no).exclude(pk=d.pk).select_related("file")
                 if d.drawing_no else [])
    siblings = access.visible(request.user, d.file.drawings.exclude(pk=d.pk)).order_by("page_no")  # 同じ図面一式の他のページ
    folder = d.file.path.rsplit("/", 1)[0] if "/" in d.file.path else ""
    bom = list(d.bom_items.all())
    uses = relations.uses(request.user, d, bom)
    found = {u["no"]: u["drawings"][0] for u in uses if u["drawings"]}
    for b in bom:
        b.ref_found = found.get(b.ref_drawing_no)
    return render(request, "drawings/detail.html", {
        "d": d, "bom": bom, "uses": uses, "used_by": relations.used_by(request.user, d), "revisions": revisions, "siblings": siblings, "q": st["q"],
        "back": back, "keep": keep, "nav": nav, "view_label": browse.VIEWS[st["v"]],
        "folder_link": browse.qs_string({"v": "folder", "f": folder}) if folder else "",
        "model_link": browse.qs_string({"v": "model", "m": f"{d.model_family}/{d.model_code}"}) if d.model_code not in ("", "_") else "",
        "series_link": browse.qs_string({"v": "series", "s": f"{d.series_prefix}/{d.series}"}) if d.series not in ("", "_") else "",
    })


def thumbnail(request, pk):
    d = get_object_or_404(access.visible(request.user), pk=pk)
    if not d.thumbnail:
        raise Http404
    p = Path(settings.MEDIA_ROOT) / d.thumbnail
    if not p.exists():
        raise Http404
    resp = FileResponse(open(p, "rb"), content_type="image/png")
    resp["Cache-Control"] = "private, max-age=3600"
    return resp


def download(request, pk):
    d = get_object_or_404(access.visible(request.user, Drawing.objects.select_related("file")), pk=pk)
    try:
        fh = open_source(d.file.path)
    except (OSError, ValueError) as e:
        return HttpResponse(f"原本を開けませんでした（{e}）。ファイルサーバーへの接続設定を確認してください。", status=502)
    AuditLog.objects.create(user=request.user, action=AuditLog.Action.DOWNLOAD, target=d.file.path)
    return FileResponse(fh, as_attachment=True, filename=d.file.filename)


def bom_xlsx(request, pk):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    d = get_object_or_404(access.visible(request.user, Drawing.objects.select_related("file")), pk=pk)
    wb = Workbook()
    ws = wb.active
    ws.title = "部品表"
    ws.append([f"{d.drawing_no} Rev.{d.revision}", d.title])
    ws.append([])
    head = ["No.", "図番又は品番", "名称・規格", "員数", "材質", "厚さ", "幅", "長さ", "関連図面", "信頼度"]
    ws.append(head)
    for c in ws[3]:
        c.font = Font(bold=True)
    low = PatternFill("solid", fgColor="FFF2CC")
    for b in d.bom_items.all():
        ws.append([b.item_no or b.row, b.part_no, b.name, float(b.qty) if b.qty is not None else None,
                   b.material, b.thickness, b.width, b.length, b.ref_drawing_no, round(b.confidence, 2)])
        if b.confidence < 0.8:
            for c in ws[ws.max_row]:
                c.fill = low
    for col, w in zip("ABCDEFGHIJ", (6, 18, 36, 7, 10, 8, 9, 9, 14, 7)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    AuditLog.objects.create(user=request.user, action=AuditLog.Action.EXPORT, target=d.file.path)
    name = f"{d.drawing_no or d.file.filename}_部品表.xlsx"
    return FileResponse(buf, as_attachment=True, filename=name)


EXPORT_MAX_DRAWINGS = 5000


def bom_export(request):
    """今の一覧（検索語・見せ方・絞り込み）に入っている図面の部品表をまとめて Excel に出す。
    型式別で型式を選んでいれば、その型式の部品集計になる。"""
    from collections import OrderedDict

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    st = browse.state_from(request.GET, request.session.get("view", "list"), access.can_see_all(request.user))
    _, qs = browse.filtered(st, request.user)
    ids = list(qs.filter(bom_items__isnull=False).distinct().order_by("drawing_no", "id")
               .values_list("id", flat=True)[:EXPORT_MAX_DRAWINGS])
    rows = (BomItem.objects.filter(drawing_id__in=ids).select_related("drawing__file")
            .order_by("drawing__drawing_no", "drawing_id", "row"))
    wb = Workbook()
    ws = wb.active
    ws.title = "部品表"
    cond = " ".join(f"{k}={v}" for k, v in st.items() if v and k != "v") or "すべて"
    ws.append([f"DWG-FIND 部品表（{browse.VIEWS[st['v']]}：{cond}）", f"図面 {len(ids)} 枚"])
    ws.append([])
    head = ["図番", "図面の名称", "型式", "ファイル", "No.", "図番又は品番", "名称・規格", "員数", "材質", "厚さ", "幅", "長さ",
            "関連図面", "信頼度"]
    ws.append(head)
    for c in ws[3]:
        c.font = Font(bold=True)
    low = PatternFill("solid", fgColor="FFF2CC")
    summary = OrderedDict()
    for b in rows:
        d = b.drawing
        qty = float(b.qty) if b.qty is not None else None
        ws.append([d.drawing_no, d.title, d.model_code if d.model_code != "_" else "", d.file.path, b.item_no or b.row, b.part_no,
                   b.name, qty, b.material, b.thickness, b.width, b.length, b.ref_drawing_no, round(b.confidence, 2)])
        if b.confidence < 0.8:
            for c in ws[ws.max_row]:
                c.fill = low
        key = b.part_no or b.name
        if key:
            s = summary.setdefault(key, {"name": b.name, "drawings": set(), "qty": 0.0, "ref": b.ref_drawing_no,
                                         "material": b.material})
            s["drawings"].add(d.drawing_no or d.file.filename)
            s["qty"] += qty or 0
    for col, w in zip("ABCDEFGHIJKLMN", (14, 24, 10, 40, 5, 18, 34, 6, 9, 7, 8, 8, 14, 7)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws2 = wb.create_sheet("集計")
    ws2.append(["図番又は品番", "名称・規格", "材質", "使っている図面の数", "員数の合計", "関連図面", "使っている図面"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for key, s in sorted(summary.items(), key=lambda kv: (-len(kv[1]["drawings"]), kv[0])):
        ws2.append([key, s["name"], s["material"], len(s["drawings"]), s["qty"] or None, s["ref"],
                    ", ".join(sorted(s["drawings"]))[:1000]])
    for col, w in zip("ABCDEFG", (20, 34, 10, 10, 10, 14, 60)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    AuditLog.objects.create(user=request.user, action=AuditLog.Action.EXPORT, target=f"部品表まとめ（{cond}）"[:1000])
    return FileResponse(buf, as_attachment=True, filename="部品表まとめ.xlsx")
