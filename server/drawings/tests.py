import io
import json
import tempfile
import time

from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.test import TestCase, override_settings

from .models import AuditLog, DictionaryEntry, Drawing, Job, SourceFile

TOKEN = "t0ken"
AUTH = {"HTTP_AUTHORIZATION": f"Bearer {TOKEN}"}
PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 32


def post_json(client, url, data, **extra):
    return client.post(url, json.dumps(data), content_type="application/json", **{**AUTH, **extra})


@override_settings(WORKER_TOKEN=TOKEN, WORKER_ALLOWED_IPS=[], ALLOWED_HOSTS=["testserver"])
class WorkerApiTests(TestCase):
    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.enterContext(override_settings(MEDIA_ROOT=self.media.name))

    def tearDown(self):
        self.media.cleanup()

    def scan(self, files):
        return post_json(self.client, "/api/internal/scan", {"files": files}).json()

    def test_token_required(self):
        self.assertEqual(self.client.get("/api/internal/health").status_code, 401)
        self.assertEqual(self.client.get("/api/internal/health", HTTP_AUTHORIZATION="Bearer wrong").status_code, 401)
        self.assertEqual(self.client.get("/api/internal/health", **AUTH).status_code, 200)

    @override_settings(WORKER_ALLOWED_IPS=["10.0.0.9"])
    def test_ip_restriction(self):
        self.assertEqual(self.client.get("/api/internal/health", **AUTH).status_code, 403)

    def test_scan_creates_jobs_only_for_new_or_changed(self):
        t = time.time()
        r = self.scan([{"path": "a/UH-1.dxf", "size": 10, "mtime": t}, {"path": "a/x.txt", "size": 1, "mtime": t}])
        self.assertEqual((r["created"], r["skipped"]), (1, 1))
        self.assertEqual(self.scan([{"path": "a/UH-1.dxf", "size": 10, "mtime": t}])["unchanged"], 1)
        self.assertEqual(Job.objects.count(), 1)  # 変化なし → 追加のジョブなし
        self.assertEqual(self.scan([{"path": "a/UH-1.dxf", "size": 11, "mtime": t}])["changed"], 1)
        self.assertEqual(Job.objects.count(), 1)  # まだ待ち行列にあるので重複させない

    def test_full_job_cycle_and_search_text(self):
        self.scan([{"path": "design/UH-3600-01_A.dxf", "size": 10, "mtime": time.time()}])
        jobs = post_json(self.client, "/api/internal/jobs/claim", {"worker": "mac", "limit": 5}).json()["jobs"]
        self.assertEqual(len(jobs), 1)
        self.assertEqual(post_json(self.client, "/api/internal/jobs/claim", {"limit": 5}).json()["jobs"], [])
        data = {"sha256": "ab" * 32,
                "drawing": {"drawing_no": "UH-3600-01", "revision": "A", "title": "ユニットハウス トイレ付", "source": "attrib", "confidence": 0.97},
                "pages": [{"page_no": 1, "text": "注記"}],
                "bom": [{"part_no": "TL-01", "name": "洋式トイレユニット", "qty": "1"}]}
        with tempfile.NamedTemporaryFile(suffix=".png") as f:
            f.write(PNG)
            f.seek(0)
            r = self.client.post(f"/api/internal/jobs/{jobs[0]['job_id']}/result", {"data": json.dumps(data), "thumbnail": f}, **AUTH)
        self.assertEqual(r.status_code, 200, r.content)
        d = Drawing.objects.get()
        self.assertEqual(d.bom_items.get().part_no, "TL-01")
        self.assertIn("TL-01", d.search_text)
        self.assertIn("ユニットハウス", d.search_text)
        self.assertTrue(d.thumbnail)
        self.assertEqual(d.file.status, SourceFile.Status.DONE)

    def test_drawing_set_pages_become_drawings(self):
        self.scan([{"path": "set/20-032 CAK-A 図面一式.pdf", "size": 10, "mtime": time.time()}])
        job = post_json(self.client, "/api/internal/jobs/claim", {"limit": 1}).json()["jobs"][0]
        data = {"sha256": "cd" * 32, "drawings": [
            {"page_no": n, "drawing": {"drawing_no": no, "title": t, "source": "ocr", "attributes": {"model": "CAK-A", "pages": 2}},
             "pages": [{"page_no": n, "text": f"注記 {no}"}], "bom": []}
            for n, no, t in [(1, "HDBY003920", "ハイキパネル/FP"), (2, "HDBY003930", "フレーム")]]}
        with tempfile.NamedTemporaryFile(suffix=".png") as a, tempfile.NamedTemporaryFile(suffix=".png") as b:
            a.write(PNG); a.seek(0); b.write(PNG); b.seek(0)
            r = self.client.post(f"/api/internal/jobs/{job['job_id']}/result",
                                 {"data": json.dumps(data), "thumb_1": a, "thumb_2": b}, **AUTH)
        self.assertEqual(r.status_code, 200, r.content)
        ds = Drawing.objects.order_by("page_no")
        self.assertEqual([(d.page_no, d.drawing_no) for d in ds], [(1, "HDBY003920"), (2, "HDBY003930")])
        self.assertTrue(all(d.thumbnail for d in ds))
        self.assertIn("CAK-A", ds[0].search_text)
        # 再解析で 1 ページになったら 2 ページ目の図面は消える
        SourceFile.objects.update(size=11)
        self.scan([{"path": "set/20-032 CAK-A 図面一式.pdf", "size": 12, "mtime": time.time()}])
        job = post_json(self.client, "/api/internal/jobs/claim", {"limit": 1}).json()["jobs"][0]
        data["drawings"] = data["drawings"][:1]
        self.client.post(f"/api/internal/jobs/{job['job_id']}/result", {"data": json.dumps(data)}, **AUTH)
        self.assertEqual(Drawing.objects.count(), 1)

    def test_nul_characters_are_removed(self):
        self.scan([{"path": "x/NUL.pdf", "size": 1, "mtime": time.time()}])
        job = post_json(self.client, "/api/internal/jobs/claim", {"limit": 1}).json()["jobs"][0]
        data = {"drawings": [{"page_no": 1, "drawing": {"drawing_no": "HB\u00000001", "title": "a\u0000b",
                                                         "attributes": {"k\u0000": "v\u0000"}},
                              "pages": [{"page_no": 1, "text": "注記\u0000本文"}], "bom": []}]}
        r = self.client.post(f"/api/internal/jobs/{job['job_id']}/result", {"data": json.dumps(data)}, **AUTH)
        self.assertEqual(r.status_code, 200, r.content)
        d = Drawing.objects.get()
        self.assertEqual((d.drawing_no, d.title), ("HB0001", "ab"))
        self.assertEqual(d.pages.get().text, "注記本文")

    def test_fail_retries_then_gives_up(self):
        self.scan([{"path": "x/UH-2.pdf", "size": 1, "mtime": time.time()}])
        for i in range(3):
            job = post_json(self.client, "/api/internal/jobs/claim", {"limit": 1}).json()["jobs"][0]
            retry = post_json(self.client, f"/api/internal/jobs/{job['job_id']}/fail", {"error": "boom"}).json()["retry"]
        self.assertFalse(retry)
        self.assertEqual(SourceFile.objects.get().status, SourceFile.Status.ERROR)

    def test_missing_only_after_complete_scan(self):
        t = time.time()
        self.scan([{"path": "a/1.dxf", "size": 1, "mtime": t}, {"path": "a/2.dxf", "size": 1, "mtime": t}])
        started = time.time() + 1
        SourceFile.objects.filter(path="a/1.dxf").update(last_seen=None)
        SourceFile.objects.filter(path="a/2.dxf").update(last_seen="2000-01-01T00:00:00Z")
        r = post_json(self.client, "/api/internal/scan/finish", {"started_at": started, "complete": False}).json()
        self.assertEqual(r["missing"], 0)
        r = post_json(self.client, "/api/internal/scan/finish", {"started_at": started, "complete": True}).json()
        self.assertEqual(r["missing"], 1)


@override_settings(ALLOWED_HOSTS=["testserver"])
class SearchViewTests(TestCase):
    def setUp(self):
        from datetime import datetime, timezone
        self.user = User.objects.create_user("u", password="pw-12345678")
        self.user.groups.add(Group.objects.get(name="User"))
        for i, (no, title, part) in enumerate([("UH-3600", "事務所", "WP-900"), ("UH-3600-01", "トイレ付事務所", "TL-01"), ("UH-3602", "コーナー柱", "UC-2400")]):
            f = SourceFile.objects.create(path=f"d/{no}.dxf", kind="dxf", size=1, mtime=datetime.now(timezone.utc))
            d = Drawing.objects.create(file=f, drawing_no=no, title=title, doc_type="drawing")
            d.bom_items.create(row=1, part_no=part, name=title)
            from .text import build_search_text
            d.search_text = build_search_text(no, title, part)
            d.save()
        for t in ["トイレ", "便所", "TOILET"]:
            DictionaryEntry.objects.create(term=t, canonical="トイレ", approved=True)
        DictionaryEntry.objects.create(term="WC", canonical="トイレ", approved=False)

    def test_login_required(self):
        self.assertEqual(self.client.get("/").status_code, 302)

    def test_search_normalizes_and_expands(self):
        self.client.login(username="u", password="pw-12345678")
        r = self.client.get("/", {"q": "ｕｈ－３６００"})
        self.assertEqual(r.context["total"], 2)
        r = self.client.get("/", {"q": "便所"})
        self.assertEqual(r.context["total"], 1)
        self.assertIn("トイレ", r.context["terms"])
        self.assertNotIn("WC", r.context["terms"])  # 未承認の語は使わない
        r = self.client.get("/", {"q": "UH-3600 トイレ"})  # AND
        self.assertEqual(r.context["total"], 1)

    def test_detail_writes_audit_log(self):
        self.client.login(username="u", password="pw-12345678")
        d = Drawing.objects.get(drawing_no="UH-3602")
        self.assertEqual(self.client.get(f"/d/{d.pk}/").status_code, 200)
        self.assertEqual(self.client.get(f"/d/{d.pk}/bom.xlsx").status_code, 200)
        self.assertEqual(AuditLog.objects.count(), 2)

    def test_back_keeps_page_and_prev_next(self):
        from datetime import datetime, timezone
        from .text import build_search_text
        for n in range(120):
            f = SourceFile.objects.create(path=f"p/ZZ{n:04d}.pdf", kind="pdf", size=1, mtime=datetime.now(timezone.utc))
            Drawing.objects.create(file=f, drawing_no=f"ZZ{n:04d}", search_text=build_search_text(f"ZZ{n:04d}"), doc_type="drawing")
        self.client.login(username="u", password="pw-12345678")
        r = self.client.get("/", {"q": "ZZ", "kind": "pdf", "page": 2})
        it = r.context["items"][0]
        d = it["d"]
        self.assertEqual((it["i"], d.drawing_no), (50, "ZZ0050"))
        self.assertIn("i=50", it["link"])
        r = self.client.get(f"/d/{d.pk}/{it['link']}")
        self.assertIn("page=2", r.context["back"])
        self.assertIn("kind=pdf", r.context["back"])
        self.assertTrue(r.context["back"].endswith(f"#f{d.file_id}"))
        nav = r.context["nav"]
        self.assertEqual((nav["pos"], nav["total"]), (51, 120))
        self.assertEqual((nav["prev"][0].drawing_no, nav["next"][0].drawing_no), ("ZZ0049", "ZZ0051"))
        self.assertIn("i=49", nav["prev"][1])
        first = Drawing.objects.get(drawing_no="ZZ0000")
        r = self.client.get(f"/d/{first.pk}/", {"q": "ZZ", "i": 0})
        self.assertNotIn("prev", r.context["nav"])
        r = self.client.get(f"/d/{first.pk}/", {"q": "ZZ", "p": 3})
        self.assertEqual(r.context["nav"], {})
        self.assertIn("page=3", r.context["back"])
        r = self.client.get("/", {"q": "ZZ"})
        self.assertEqual(r.context["window"], [1, 2, 3])


class ClassifyTests(TestCase):
    def test_series(self):
        from .classify import series_of
        self.assertEqual(series_of("HB0011XXXX"), ("HB", "HB0011"))
        self.assertEqual(series_of("HDBY003920"), ("HDBY", "HDBY0039"))
        self.assertEqual(series_of("HDBY003930"), ("HDBY", "HDBY0039"))
        self.assertEqual(series_of("UH-3600-01"), ("UH", "UH-3600"))
        self.assertEqual(series_of("UH-3600"), ("UH", "UH-3600"))
        self.assertEqual(series_of("1234567890"), ("数字", "123456"))
        self.assertEqual(series_of(""), ("_", "_"))

    def test_model(self):
        from .classify import model_of
        self.assertEqual(model_of("図面/x/20-032 CAK-A 多用途ﾊｳｽ 図面一式.pdf", {"model": "CAK-A"}), ("CAK", "CAK-A"))
        self.assertEqual(model_of("図面/dxfCAK-40A 田の字36R-221003.dxf"), ("CAK", "CAK-40A"))
        self.assertEqual(model_of("図面 DXF・DWG・JW・PDF/JH/JH 標準/HB0011XXXX 差替3.pdf"), ("JH", "JH"))
        self.assertEqual(model_of("図面/CAG-12/部品/UH-3600-01.dxf"), ("CAG", "CAG-12"))
        self.assertEqual(model_of("図面/部品/UH-3600-01.dxf"), ("_", "_"))  # 図番は型式にしない
        self.assertEqual(model_of("PDF/DXF-1/a.pdf"), ("_", "_"))
        # OCR で空白だけ・文字以外が入っていても止まらない
        for bad in (" ", "\u3000", 123, ["CAK"], None, ""):
            self.assertEqual(model_of("a/b.pdf", {"model": bad}), ("_", "_"))


@override_settings(ALLOWED_HOSTS=["testserver"])
class BrowseViewTests(TestCase):
    def setUp(self):
        from datetime import datetime, timezone
        from . import classify
        from .text import build_search_text
        User.objects.create_user("u", password="pw-12345678").groups.add(Group.objects.get(name="User"))
        self.client.login(username="u", password="pw-12345678")
        now = datetime.now(timezone.utc)
        rows = [("図面/CAK/CAK-A/20-032 CAK-A 図面一式.pdf", [(1, "HDBY003920"), (2, "HDBY003930"), (3, "HDBY004010")]),
                ("図面/CAK/dxfCAK-40A 田の字36R-221003.dxf", [(1, "")]),
                ("図面/JH/HB0011XXXX差替3 スペーサー.pdf", [(1, "HB0011XXXX")]),
                ("資料/申請書.pdf", [(1, "")])]
        for path, pages in rows:
            f = SourceFile.objects.create(path=path, kind=path[-3:], size=1, mtime=now)
            for page_no, no in pages:
                d = Drawing(file=f, page_no=page_no, drawing_no=no, attributes={"pages": len(pages)},
                            search_text=build_search_text(no, path), confidence=0.75)
                classify.apply(d, text="")
                d.doc_type = "drawing"
                d.save()

    def test_list_collapses_drawing_sets(self):
        r = self.client.get("/", {"v": "list"})
        self.assertEqual(r.context["total"], 4)  # 図面一式の 3 枚は 1 件
        set_item = next(it for it in r.context["items"] if it["size"] == 3)
        self.assertEqual(set_item["d"].page_no, 1)
        self.assertContains(r, "一式 3枚")
        self.assertEqual(r.context["items"][-1]["d"].drawing_no, "")  # 図番のないものは後ろ
        # 検索で 2 ページ目だけ一致したら、代表は 2 ページ目
        r = self.client.get("/", {"q": "HDBY003930"})
        self.assertEqual([it["d"].page_no for it in r.context["items"]], [2])

    def test_folder_tree(self):
        r = self.client.get("/", {"v": "folder"})
        self.assertEqual([(n["label"], n["n"]) for n in r.context["nodes"]], [("図面", 5), ("資料", 1)])
        r = self.client.get("/", {"v": "folder", "f": "図面/CAK"})
        labels = [(n["label"], n["depth"], n["current"]) for n in r.context["nodes"]]
        self.assertIn(("CAK", 1, True), labels)
        self.assertIn(("CAK-A", 2, False), labels)
        self.assertIn(("JH", 1, False), labels)
        self.assertIn(("資料", 0, False), labels)
        self.assertEqual(r.context["total"], 2)  # CAK 以下のファイル
        self.assertEqual([c[0] for c in r.context["crumbs"]], ["図面", "CAK"])
        # 見せ方を覚える
        r = self.client.get("/")
        self.assertEqual(r.context["st"]["v"], "folder")

    def test_model_and_series(self):
        r = self.client.get("/", {"v": "model"})
        self.assertEqual({n["label"]: n["n"] for n in r.context["nodes"]}, {"CAK": 4, "JH": 1, "（型式なし）": 1})
        r = self.client.get("/", {"v": "model", "m": "CAK"})
        self.assertIn("CAK-40A", [n["label"] for n in r.context["nodes"]])
        r = self.client.get("/", {"v": "model", "m": "CAK/CAK-40A"})
        self.assertEqual(r.context["total"], 1)
        labels = [n["label"] for n in r.context["nodes"]]
        self.assertIn("CAK-A", labels)  # 選んだ型式の兄弟も見える
        self.assertIn("JH", labels)
        r = self.client.get("/", {"v": "series", "s": "HDBY/HDBY0039"})
        self.assertEqual(r.context["total"], 1)
        self.assertEqual(r.context["drawing_total"], 2)

    def test_detail_nav_in_folder_view(self):
        r = self.client.get("/", {"v": "folder", "f": "図面"})
        it = r.context["items"][0]
        r = self.client.get(f"/d/{it['d'].pk}/{it['link']}")
        self.assertIn("v=folder", r.context["back"])
        self.assertIn("f=", r.context["back"])
        self.assertEqual(r.context["nav"]["total"], 3)
        self.assertTrue(r.context["folder_link"])


class DocTypeTests(TestCase):
    def test_rules(self):
        from .classify import doc_type_of as t
        self.assertEqual(t("物件/2024/山田様 契約書.pdf"), "contract")
        self.assertEqual(t("物件/見積/HB0011XXXX.pdf", "HB0011XXXX", "filename", 0.9), "contract")  # フォルダの語が先
        self.assertEqual(t("物件/A邸/公図.pdf"), "site")
        self.assertEqual(t("物件/A邸/建築確認申請 図書一式.pdf"), "application")
        self.assertEqual(t("図面/HB0011XXXX差替3 スペーサー.pdf", "HB0011XXXX", "filename", 0.9), "drawing")
        self.assertEqual(t("図面/20-032 CAK-A 多用途ﾊｳｽ 図面一式.pdf"), "drawing")
        self.assertEqual(t("図面/dxfCAK-40A 田の字36R-221003.dxf", kind="dxf"), "drawing")
        self.assertEqual(t("図面/scan001.pdf", "HDBY003920", "ocr", 0.75), "drawing")
        # 名前で分からなければ本文
        self.assertEqual(t("資料/scan002.pdf", text="工事請負契約書 契約金額 金 1,000,000 円 収入印紙"), "contract")
        self.assertEqual(t("資料/scan003.pdf", text="地番 123-4 地積 250.00m2 公図 写し"), "site")
        self.assertEqual(t("資料/scan004.pdf", text="図番 HB0011 尺度 1/10 材質 SS400"), "drawing")
        self.assertEqual(t("資料/scan005.pdf", text="お知らせ"), "other")
        # 「確認申請」の中の「申請」は別の語として数えない（1 語だけでは申請書類にしない）
        self.assertEqual(t("資料/scan006.pdf", text="確認申請"), "other")
        self.assertEqual(t("資料/scan007.pdf", text="契約書"), "other")
        # ファイル名はフォルダ名より優先
        self.assertEqual(t("物件/見積/公図.pdf"), "site")
        from .classify import explain
        self.assertEqual(explain("物件/見積/a.pdf")[1], "フォルダ名「見積」（見積）")  # 迷ったら other（ゲストに見せない）

    def test_real_cases_2026_09_30(self):
        """ユーザー確認で書類だった 6 件（図面は 1 件）。複数ページ・書類の多いフォルダ・書類名の語。"""
        from .classify import doc_type_of as t
        base = "図面原紙･資料 PDF/物件対応ファイル他/"
        multi = {"pages": 5}
        # 複数ページで、図面一式として分かれていない PDF → 書類
        self.assertEqual(t(base + "2004年･H16 新事業体形成の為の要素開発/20241127081623360.pdf", "HB0011", "ocr", 0.75, "pdf",
                           "工程表 4月 5月", multi), "general")
        self.assertEqual(t(base + "2003年・H15 宝くじBOX販売物件/20250716081819252.pdf", "HDBY003920", "ocr", 0.75, "pdf",
                           "発注伝票 品名 数量", multi), "contract")
        self.assertEqual(t(base + "1993年 試験研究複合ファイル/x.pdf", "HB0011", "ocr", 0.75, "pdf", "", multi), "other")
        # 1 ページでも書類の多いフォルダでは、図番らしい文字だけでは図面にしない
        self.assertEqual(t(base + "2002年・H14 築地銀だこ/20250611082900706.pdf", "HDBY003920", "ocr", 0.75, "pdf",
                           "部品寸法表 A 100 B 200"), "general")
        self.assertEqual(t(base + "2000年・H12 大型仮設資料/20250625083031386.pdf", "HB0011", "ocr", 0.75, "pdf", "凛議書"), "general")
        self.assertEqual(t(base + "1993年 試験研究複合ファイル/20241204080449608.pdf", "HB0011", "ocr", 0.75, "pdf", "理由書"), "general")
        self.assertEqual(t(base + "a/20241204080525522.pdf", "HB0011", "ocr", 0.75, "pdf", "申請書 住所"), "application")
        self.assertEqual(t(base + "a/20241204080525522.pdf", "HB0011", "ocr", 0.75, "pdf", ""), "other")
        # 書類の多いフォルダでも、図番と表題欄の語がそろえば図面
        self.assertEqual(t(base + "a/20250528083819213.pdf", "HDBY003920", "ocr", 0.75, "pdf", "図番 HDBY003920 尺度 1/10 材質 SS400"),
                         "drawing")
        # 図面一式としてページごとに分かれたもの（attributes に page がある）は今まで通り図面
        self.assertEqual(t("図面/20-032 x.pdf", "HDBY003920", "ocr", 0.75, "pdf", "", {"pages": 12, "page": 3}), "drawing")
        # 部品図のフォルダの 1 ページ PDF は今まで通り
        self.assertEqual(t("図面 DXF・DWG・JW・PDF/CAK/HB0011XXXX差替3 スペーサー.pdf", "HB0011XXXX", "filename", 0.9, "pdf"), "drawing")

    def test_manual_fix_is_kept(self):
        from datetime import datetime, timezone
        from . import classify
        f = SourceFile.objects.create(path="資料/公図.pdf", kind="pdf", size=1, mtime=datetime.now(timezone.utc))
        d = Drawing.objects.create(file=f, doc_type="drawing", doc_type_fixed=True)
        classify.apply(d, text="")
        self.assertEqual(d.doc_type, "drawing")
        d.doc_type_fixed = False
        classify.apply(d, text="")
        self.assertEqual(d.doc_type, "site")


@override_settings(ALLOWED_HOSTS=["testserver"])
class AccessTests(TestCase):
    def setUp(self):
        from datetime import datetime, timezone
        from .text import build_search_text
        self.media = tempfile.TemporaryDirectory()
        self.enterContext(override_settings(MEDIA_ROOT=self.media.name))
        now = datetime.now(timezone.utc)
        self.d = {}
        for key, path, doc_type in [("prod", "図面/HB0011XXXX スペーサー.pdf", "drawing"),
                                    ("contract", "物件/山田様 契約書.pdf", "contract"),
                                    ("site", "物件/公図.pdf", "site")]:
            f = SourceFile.objects.create(path=path, kind="pdf", size=1, mtime=now)
            import pathlib
            (pathlib.Path(self.media.name) / f"{key}.png").write_bytes(PNG)
            self.d[key] = Drawing.objects.create(file=f, drawing_no="HB0011XXXX" if key == "prod" else "", doc_type=doc_type,
                                                 thumbnail=f"{key}.png", search_text=build_search_text(path, "共通語"))
        self.users = {}
        for name, group in [("admin1", "Admin"), ("user1", "User"), ("guest1", "Guest"), ("nogroup", None)]:
            u = User.objects.create_user(name, password="pw-12345678")
            if group:
                u.groups.add(Group.objects.get(name=group))
            self.users[name] = u

    def login(self, name):
        self.client.logout()
        self.client.login(username=name, password="pw-12345678")

    def test_guest_sees_only_product_drawings(self):
        for name in ("guest1", "nogroup"):  # グループなしもゲスト扱い
            self.login(name)
            r = self.client.get("/", {"q": "共通語", "t": "all"})  # t=all を付けても無視
            self.assertEqual(r.context["total"], 1)
            self.assertNotContains(r, "契約書")
            self.assertNotContains(r, 'name="t"')  # 種別の切替は出さない
            r = self.client.get("/", {"v": "folder", "t": "all"})
            self.assertEqual([n["label"] for n in r.context["nodes"]], ["図面"])  # 階層の件数にも出ない
            for key in ("contract", "site"):
                pk = self.d[key].pk
                for url in (f"/d/{pk}/", f"/d/{pk}/thumb.png", f"/d/{pk}/download", f"/d/{pk}/bom.xlsx"):
                    self.assertEqual(self.client.get(url).status_code, 404, url)
            self.assertEqual(self.client.get(f"/d/{self.d['prod'].pk}/").status_code, 200)
            self.assertEqual(self.client.get(f"/d/{self.d['prod'].pk}/thumb.png").status_code, 200)
            self.assertNotEqual(self.client.get("/admin/").status_code, 200)

    def test_user_sees_all_but_no_admin(self):
        self.login("user1")
        r = self.client.get("/", {"q": "共通語"})
        self.assertEqual(r.context["total"], 1)  # 既定は製品図面だけ
        r = self.client.get("/", {"q": "共通語", "t": "all"})
        self.assertEqual(r.context["total"], 3)
        r = self.client.get("/", {"q": "共通語", "t": "contract"})
        self.assertEqual(r.context["total"], 1)
        self.assertEqual(self.client.get(f"/d/{self.d['contract'].pk}/").status_code, 200)
        self.assertEqual(self.client.get(f"/d/{self.d['contract'].pk}/thumb.png").status_code, 200)
        self.assertEqual(self.client.get("/admin/").status_code, 302)  # 管理画面はログイン画面へ

    def test_admin_group_gets_admin_site(self):
        u = User.objects.get(username="admin1")
        self.assertTrue(u.is_staff and u.is_superuser)
        self.login("admin1")
        self.assertEqual(self.client.get("/admin/").status_code, 200)
        self.assertEqual(self.client.get(f"/d/{self.d['site'].pk}/").status_code, 200)
        # Admin から外すと管理権限も外れる（他に管理者がいるとき）
        User.objects.create_superuser("root2", password="pw-12345678")
        u.groups.remove(Group.objects.get(name="Admin"))
        u.refresh_from_db()
        self.assertFalse(u.is_staff or u.is_superuser)

    def test_last_admin_is_not_demoted(self):
        u = User.objects.get(username="admin1")
        u.groups.remove(Group.objects.get(name="Admin"))
        u.refresh_from_db()
        self.assertTrue(u.is_superuser)

    def test_superuser_joins_admin_group(self):
        u = User.objects.create_superuser("root3", password="pw-12345678")
        self.assertTrue(u.groups.filter(name="Admin").exists())

    def test_set_group_command(self):
        import io
        out = io.StringIO()
        call_command("set_group", "guest1", "user", stdout=out)
        u = User.objects.get(username="guest1")
        self.assertEqual(sorted(u.groups.values_list("name", flat=True)), ["User"])
        call_command("set_group", "--list", stdout=out)
        self.assertIn("guest1", out.getvalue())


@override_settings(ALLOWED_HOSTS=["testserver"])
class LoginLogoutTests(TestCase):
    def test_logout_goes_to_normal_login(self):
        User.objects.create_user("u", password="pw-12345678")
        self.client.login(username="u", password="pw-12345678")
        r = self.client.post("/accounts/logout/")
        self.assertRedirects(r, "/accounts/login/", fetch_redirect_response=False)

    def test_admin_logout_and_login_use_normal_login(self):
        User.objects.create_superuser("root", password="pw-12345678")
        self.client.login(username="root", password="pw-12345678")
        r = self.client.post("/admin/logout/")
        self.assertRedirects(r, "/accounts/login/", fetch_redirect_response=False)
        r = self.client.get("/admin/", follow=True)
        self.assertEqual(r.redirect_chain[-1][0], "/accounts/login/?next=%2Fadmin%2F")
        self.assertTemplateUsed(r, "drawings/login.html")
        r = self.client.post("/accounts/login/?next=/admin/", {"username": "root", "password": "pw-12345678"})
        self.assertRedirects(r, "/admin/", fetch_redirect_response=False)


class ReclassifyCommandTests(TestCase):
    def test_dry_run_does_not_write(self):
        import io
        from datetime import datetime, timezone
        f = SourceFile.objects.create(path="物件対応ファイル他/a/20241127081623360.pdf", kind="pdf", size=1,
                                      mtime=datetime.now(timezone.utc))
        d = Drawing.objects.create(file=f, drawing_no="HB0011", confidence=0.75, source="ocr", doc_type="drawing",
                                   attributes={"pages": 3})
        out = io.StringIO()
        call_command("reclassify", "--dry-run", stdout=out)
        self.assertIn("drawing → other", out.getvalue())
        d.refresh_from_db()
        self.assertEqual(d.doc_type, "drawing")
        call_command("reclassify", stdout=io.StringIO())
        d.refresh_from_db()
        self.assertEqual(d.doc_type, "other")


@override_settings(WORKER_TOKEN=TOKEN, WORKER_ALLOWED_IPS=[], ALLOWED_HOSTS=["testserver"])
class BomRelationTests(TestCase):
    """部品表の登録・関連図面（使う部品／使っている図面）・まとめ Excel・入れ直し。"""

    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.enterContext(override_settings(MEDIA_ROOT=self.media.name))
        u = User.objects.create_user("u", password="pw-12345678")
        u.groups.add(Group.objects.get(name="User"))
        self.client.login(username="u", password="pw-12345678")

    def post_result(self, path, no, bom, note_refs=None, kind="dxf"):
        post_json(self.client, "/api/internal/scan", {"files": [{"path": path, "size": 1, "mtime": time.time()}]})
        job = post_json(self.client, "/api/internal/jobs/claim", {"limit": 1}).json()["jobs"][0]
        attrs = {"note_refs": note_refs} if note_refs else {}
        data = {"drawings": [{"page_no": 1, "drawing": {"drawing_no": no, "title": no + " の図", "source": "attrib",
                                                          "confidence": 0.97, "attributes": attrs},
                              "pages": [{"page_no": 1, "text": no}], "bom": bom}]}
        r = self.client.post(f"/api/internal/jobs/{job['job_id']}/result", {"data": json.dumps(data)}, HTTP_AUTHORIZATION=f"Bearer {TOKEN}")
        self.assertEqual(r.status_code, 200, r.content)
        return Drawing.objects.get(file__path=path)

    def test_relations_and_export(self):
        panel = self.post_result("図面/CAK/HDBY003920.dxf", "HDBY003920", [
            {"item_no": "1", "part_no": "HDBY003930", "name": "フレーム/ハイキパネル", "qty": "1", "ref_drawing_no": "HDBY003930",
             "confidence": 0.8},
            {"item_no": "2", "part_no": "P", "name": "エンボスカラー", "qty": "1", "material": "C01", "thickness": "0.25",
             "width": "914.00", "length": "2443", "confidence": 1.0},
        ], note_refs=["HUCP101040"])
        frame = self.post_result("図面/CAK/HDBY003930.dxf", "HDBY003930", [])
        b = panel.bom_items.get(row=2)
        self.assertEqual((b.thickness, b.width, b.length), ("0.25", "914.00", "2443"))
        # パネル → 使う部品にフレーム、注記の参照（図面はまだない）
        r = self.client.get(f"/d/{panel.pk}/")
        uses = {u["no"]: u for u in r.context["uses"]}
        self.assertEqual([x.pk for x in uses["HDBY003930"]["drawings"]], [frame.pk])
        self.assertEqual(uses["HUCP101040"]["drawings"], [])
        self.assertEqual(r.context["bom"][0].ref_found.pk, frame.pk)
        # フレーム → 使っている図面にパネル
        r = self.client.get(f"/d/{frame.pk}/")
        self.assertEqual([u["drawing"].pk for u in r.context["used_by"]], [panel.pk])
        # 注記で参照された図面が後から登録されても、すぐにつながる
        detail = self.post_result("図面/HUCP101040.dxf", "HUCP101040", [])
        r = self.client.get(f"/d/{detail.pk}/")
        self.assertEqual([(u["drawing"].pk, u["via"]) for u in r.context["used_by"]], [(panel.pk, "注記")])
        # 品番で逆引き：部品表の品番で検索すると、使っている図面が出る
        r = self.client.get("/", {"q": "HDBY003930"})
        self.assertIn(panel.pk, [it["d"].pk for it in r.context["items"]])
        # まとめ Excel（一覧の条件どおり）
        from openpyxl import load_workbook
        r = self.client.get("/export/bom.xlsx", {"q": "HDBY"})
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(b"".join(r.streaming_content)))
        rows = list(wb["部品表"].iter_rows(min_row=4, values_only=True))
        self.assertEqual([x[5] for x in rows], ["HDBY003930", "P"])
        summary = list(wb["集計"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary[0][0], "HDBY003930")
        self.assertTrue(AuditLog.objects.filter(action="export", target__startswith="部品表まとめ").exists())

    def test_assembly_refs_link_both_ways(self):
        parent = self.post_result("図面/HDZY007950.dxf", "HDZY007950", [])
        part = self.post_result("図面/HDZY007970.dxf", "HDZY007970", [])
        Drawing.objects.filter(pk=part.pk).update(attributes={"assembly_refs": ["HDZY007950"]})
        r = self.client.get(f"/d/{part.pk}/")
        self.assertEqual([(u["drawing"].pk, u["via"]) for u in r.context["used_by"]], [(parent.pk, "組立図番")])
        r = self.client.get(f"/d/{parent.pk}/")
        self.assertEqual([(u["drawings"][0].pk, u["via"]) for u in r.context["uses"]], [(part.pk, "組立図番")])

    def test_guest_does_not_see_hidden_relations(self):
        panel = self.post_result("物件/山田様 契約書.dxf", "HDBY003920",
                                 [{"item_no": "1", "part_no": "HDBY003930", "ref_drawing_no": "HDBY003930"}])
        frame = self.post_result("図面/HDBY003930.dxf", "HDBY003930", [])
        self.assertEqual(Drawing.objects.get(pk=panel.pk).doc_type, "contract")
        g = User.objects.create_user("g", password="pw-12345678")
        g.groups.add(Group.objects.get(name="Guest"))
        self.client.logout()
        self.client.login(username="g", password="pw-12345678")
        r = self.client.get(f"/d/{frame.pk}/")
        self.assertEqual(r.context["used_by"], [])  # 契約書の中の部品表からはつながない
        r = self.client.get("/export/bom.xlsx", {"t": "all"})
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(b"".join(r.streaming_content)))
        self.assertEqual(list(wb["部品表"].iter_rows(min_row=4, values_only=True)), [])

    def test_requeue_command(self):
        from datetime import datetime, timezone
        now = datetime.now(timezone.utc)
        for p in ("a/1.dxf", "a/2.dxf", "b/3.dxf", "a/4.pdf"):
            SourceFile.objects.create(path=p, kind=p[-3:], size=1, mtime=now, status="done")
        out = io.StringIO()
        call_command("requeue", "--kind", "dxf", "--dry-run", stdout=out)
        self.assertIn("3 件", out.getvalue())
        self.assertEqual(Job.objects.count(), 0)
        call_command("requeue", "--kind", "dxf", "--path", "a", stdout=io.StringIO())
        self.assertEqual(sorted(Job.objects.values_list("file__path", "priority")), [("a/1.dxf", 200), ("a/2.dxf", 200)])
        call_command("requeue", "--kind", "dxf", stdout=io.StringIO())  # 待ち中のものは重ねない
        self.assertEqual(Job.objects.count(), 3)


def _dxf_bytes(frame=(420, 297), scale_factor=1.0, text="HB0011XXXX 部品図"):
    """用紙の大きさの外枠と内枠、100mm の線を描いた DXF（紙の上の mm。scale_factor 倍で実寸の図面にする）。"""
    import ezdxf

    doc = ezdxf.new("R2010")
    msp = doc.modelspace()
    w, h = frame
    k = scale_factor
    msp.add_lwpolyline([(0, 0), (w * k, 0), (w * k, h * k), (0, h * k)], close=True)
    msp.add_lwpolyline([(15 * k, 10 * k), ((w - 10) * k, 10 * k), ((w - 10) * k, (h - 10) * k), (15 * k, (h - 10) * k)],
                       close=True)
    msp.add_line((50 * k, 50 * k), (150 * k, 50 * k))
    msp.add_text(text, height=5 * k).set_placement((60 * k, 60 * k))
    msp.add_line((-30 * k, (h + 20) * k), (-10 * k, (h + 20) * k))  # 図枠の外のメモ（用紙には入らない）
    s = io.StringIO()
    doc.write(s)
    return s.getvalue().encode("utf-8")


class DxfPdfTests(TestCase):
    def _pdf(self, data, scale=""):
        import pymupdf
        from . import dxfpdf

        pdf, sheet = dxfpdf.render(data, scale)
        page = pymupdf.open(stream=pdf, filetype="pdf")[0]
        return page, sheet

    @staticmethod
    def _mm(pt):
        return round(pt / 72 * 25.4)

    def _has_line_mm(self, page, mm):
        for p in page.get_drawings():
            for it in p["items"]:
                if it[0] == "l" and abs(it[1].y - it[2].y) < 0.2 and abs(abs(it[1].x - it[2].x) / 72 * 25.4 - mm) < 0.8:
                    return True
        return False

    def test_scale_parse(self):
        from .dxfpdf import scale_denominator
        self.assertEqual(scale_denominator("1/30"), 30)
        self.assertEqual(scale_denominator("S=1:25"), 25)
        self.assertEqual(scale_denominator("１／１０"), 10)
        self.assertIsNone(scale_denominator("NTS"))

    def test_a3_drawing_is_a3_full_size(self):
        page, sheet = self._pdf(_dxf_bytes((420, 297)), "1/10")
        self.assertEqual((sheet.paper, sheet.exact), ("A3", True))
        self.assertEqual((self._mm(page.rect.width), self._mm(page.rect.height)), (420, 297))
        self.assertTrue(self._has_line_mm(page, 100))  # 100mm の線が 100mm で出る（原寸）

    def test_a2_drawing_stays_a2(self):
        page, sheet = self._pdf(_dxf_bytes((594, 420)), "1/30")
        self.assertEqual(sheet.paper, "A2")
        self.assertEqual((self._mm(page.rect.width), self._mm(page.rect.height)), (594, 420))
        self.assertTrue(self._has_line_mm(page, 100))

    def test_portrait_a3(self):
        page, sheet = self._pdf(_dxf_bytes((297, 420)))
        self.assertEqual(sheet.paper, "A3")
        self.assertEqual((self._mm(page.rect.width), self._mm(page.rect.height)), (297, 420))

    def test_real_size_model_uses_title_scale(self):
        page, sheet = self._pdf(_dxf_bytes((420, 297), scale_factor=30), "1/30")
        self.assertEqual((sheet.paper, sheet.factor), ("A3", 30))
        self.assertTrue(self._has_line_mm(page, 100))  # 3000mm の線が 1/30 で 100mm


@override_settings(ALLOWED_HOSTS=["testserver"])
class PrintPdfViewTests(TestCase):
    def setUp(self):
        import pathlib
        from datetime import datetime, timezone
        self.media = tempfile.TemporaryDirectory()
        self.src = tempfile.TemporaryDirectory()
        self.enterContext(override_settings(MEDIA_ROOT=self.media.name, SOURCE_ROOT=self.src.name))
        (pathlib.Path(self.src.name) / "図面").mkdir()
        (pathlib.Path(self.src.name) / "図面" / "HB0011XXXX.dxf").write_bytes(_dxf_bytes())
        now = datetime.now(timezone.utc)
        f = SourceFile.objects.create(path="図面/HB0011XXXX.dxf", kind="dxf", size=1, mtime=now)
        self.d = Drawing.objects.create(file=f, drawing_no="HB0011XXXX", doc_type="drawing", scale="1/10")
        pf = SourceFile.objects.create(path="図面/scan.pdf", kind="pdf", size=1, mtime=now)
        self.p = Drawing.objects.create(file=pf, drawing_no="HB0012XXXX", doc_type="drawing")
        c = SourceFile.objects.create(path="物件/契約.dxf", kind="dxf", size=1, mtime=now)
        self.c = Drawing.objects.create(file=c, doc_type="contract")
        u = User.objects.create_user("g", password="pw-12345678")
        u.groups.add(Group.objects.get(name="Guest"))
        self.client.login(username="g", password="pw-12345678")

    def test_download_and_cache(self):
        r = self.client.get(f"/d/{self.d.pk}/")
        self.assertContains(r, "印刷用PDF")
        r = self.client.get(f"/d/{self.d.pk}/print.pdf")
        self.assertEqual(r.status_code, 200)
        body = b"".join(r.streaming_content)
        self.assertTrue(body.startswith(b"%PDF"))
        self.assertIn("HB0011XXXX_A3.pdf", r["Content-Disposition"].encode("latin-1").decode("utf-8", "replace")
                      if "filename*" not in r["Content-Disposition"] else r["Content-Disposition"])
        self.assertTrue(AuditLog.objects.filter(target__contains="PDF A3").exists())
        import pathlib
        self.assertEqual(len(list((pathlib.Path(self.media.name) / "print").glob("*.pdf"))), 1)
        (pathlib.Path(self.src.name) / "図面" / "HB0011XXXX.dxf").unlink()  # 2 回目は保存した PDF（原本を読まない）
        self.assertEqual(self.client.get(f"/d/{self.d.pk}/print.pdf").status_code, 200)

    def test_not_for_pdf_or_hidden(self):
        self.assertNotContains(self.client.get(f"/d/{self.p.pk}/"), "印刷用PDF")
        self.assertEqual(self.client.get(f"/d/{self.p.pk}/print.pdf").status_code, 404)
        self.assertEqual(self.client.get(f"/d/{self.c.pk}/print.pdf").status_code, 404)  # ゲストに見えない書類


class BomReportCommandTests(TestCase):
    def test_report_and_worst(self):
        from datetime import datetime, timezone
        from drawings.models import BomItem
        now = datetime.now(timezone.utc)
        f = SourceFile.objects.create(path="図面/a.pdf", kind="pdf", size=1, mtime=now)
        d = Drawing.objects.create(file=f, drawing_no="HUXP010143", doc_type="drawing")
        Drawing.objects.create(file=SourceFile.objects.create(path="図面/b.pdf", kind="pdf", size=1, mtime=now),
                               drawing_no="HUXP010153")
        for i, (p, ref) in enumerate([("HUXP010153", "HUXP010153"), ("1UXP010310", ""), ("HU99003.060", "")]):
            BomItem.objects.create(drawing=d, row=i + 1, part_no=p, name="パネル", ref_drawing_no=ref)
        out = io.StringIO()
        call_command("bom_report", stdout=out)
        self.assertIn("部品表あり 1 件 / 行 3", out.getvalue())
        self.assertIn("うち図面が見つかる 1（100%）", out.getvalue())
        out = io.StringIO()
        call_command("bom_report", worst=5, stdout=out)
        self.assertIn("HUXP010143,3,0.67", out.getvalue())
